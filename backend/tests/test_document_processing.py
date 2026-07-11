import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from app.services.document_service import (
    DocumentServiceError,
    _clean_text,
    _extract_docx_text,
    _extract_text,
    _extract_xlsx_text,
    _remove_repeated_page_margins,
)
from app.services.rag_service import _lexical_score


class DocumentProcessingTests(unittest.TestCase):
    def test_clean_text_preserves_table_columns_and_removes_controls(self):
        text = " 名称\t  数量  \x00\n\u200b商品 A\t  2 "
        self.assertEqual(_clean_text(text), "名称\t数量\n商品 A\t2")

    def test_repeated_pdf_headers_and_footers_are_removed(self):
        pages = [
            (1, "企业手册\n第一章正文\n第 1 页"),
            (2, "企业手册\n第二章正文\n第 2 页"),
            (3, "企业手册\n第三章正文\n第 3 页"),
        ]
        cleaned = _remove_repeated_page_margins(pages)
        combined = "\n".join(text for _, text in cleaned)
        self.assertNotIn("企业手册", combined)
        self.assertNotIn("第 1 页", combined)
        self.assertIn("第一章正文", combined)
        self.assertIn("第三章正文", combined)

    def test_docx_preserves_paragraph_and_table_order(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.docx"
            document = Document()
            document.add_paragraph("段落一")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "产品"
            table.cell(0, 1).text = "价格"
            document.add_paragraph("段落二")
            document.save(path)
            text = _extract_docx_text(path)
        self.assertLess(text.index("段落一"), text.index("[table 1]"))
        self.assertLess(text.index("[table 1]"), text.index("段落二"))
        self.assertIn("产品\t价格", text)

    def test_xlsx_preserves_sheet_and_columns(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "订单"
            worksheet.append(["订单号", "金额"])
            worksheet.append(["10001", 88.5])
            workbook.save(path)
            text = _extract_xlsx_text(path)
        self.assertIn("[sheet 订单]", text)
        self.assertIn("订单号\t金额", text)
        self.assertIn("10001\t88.5", text)

    def test_invalid_office_pointer_has_clear_error(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "sample.xlsx"
            path.write_text("version https://git-lfs.github.com/spec/v1", encoding="utf-8")
            with self.assertRaisesRegex(DocumentServiceError, "Git LFS"):
                _extract_text(path)

    def test_non_utf8_chinese_text_is_detected(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "policy.txt"
            expected = "企业退款政策：收货后七天内可申请退款。" * 20
            path.write_bytes(expected.encode("gb18030"))
            actual = _extract_text(path)
        self.assertEqual(actual, _clean_text(expected))

    def test_lexical_score_rejects_generic_overlap(self):
        unrelated = _lexical_score("公司的退款政策是什么", "公司的办公时间是上午九点")
        exact = _lexical_score("退款政策", "公司的退款政策为七天内可申请")
        self.assertLess(unrelated, 0.5)
        self.assertEqual(exact, 1.0)


if __name__ == "__main__":
    unittest.main()
