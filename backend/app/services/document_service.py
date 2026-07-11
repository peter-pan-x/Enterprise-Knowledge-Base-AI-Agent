import asyncio
import json
import os
import re
import shutil
import sqlite3
import tempfile
import threading
import unicodedata
from collections import Counter
from datetime import date, datetime
from datetime import UTC
from pathlib import Path
from uuid import uuid4

from fastapi import UploadFile
from charset_normalizer import from_bytes
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
import pymupdf
from pypdf import PdfReader
from rapidocr_onnxruntime import RapidOCR
import xlrd

from app.schemas.document import DocumentIndexStatus, DocumentRecord
from app.schemas.knowledge_base import KnowledgeBaseCreate, KnowledgeBaseRecord, KnowledgeBaseUpdate
from app.services.auth_service import DATABASE_PATH

BASE_DIR = Path(__file__).resolve().parents[2]
DATA_DIR = BASE_DIR / "data"
UPLOAD_DIR = DATA_DIR / "uploads"
TEXT_DIR = DATA_DIR / "texts"
METADATA_PATH = DATA_DIR / "documents.json"

ALLOWED_EXTENSIONS = {
    ".pdf", ".txt", ".md", ".markdown", ".csv",
    ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff",
    ".docx", ".xlsx", ".xls",
}
MAX_UPLOAD_SIZE_BYTES = 15 * 1024 * 1024
_METADATA_LOCK = threading.RLock()
_OCR_LOCK = threading.Lock()
_OCR_ENGINE: RapidOCR | None = None
MAX_OCR_PAGES = 50


class DocumentServiceError(Exception):
    pass


def _ensure_storage() -> None:
    with _METADATA_LOCK:
        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        TEXT_DIR.mkdir(parents=True, exist_ok=True)
        initialize_document_storage()


def initialize_document_storage() -> None:
    DATABASE_PATH.parent.mkdir(parents=True, exist_ok=True)
    with _document_connection() as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS documents (
                id TEXT PRIMARY KEY,
                filename TEXT NOT NULL,
                content_type TEXT NOT NULL,
                status TEXT NOT NULL,
                size_bytes INTEGER NOT NULL,
                text_length INTEGER NOT NULL,
                created_at TEXT NOT NULL,
                error_message TEXT,
                index_status TEXT NOT NULL,
                indexed_chunks INTEGER NOT NULL,
                index_error_message TEXT,
                is_enabled INTEGER NOT NULL DEFAULT 1,
                knowledge_base_id TEXT NOT NULL DEFAULT 'default',
                category TEXT NOT NULL DEFAULT 'general'
            )
            """
        )
        columns = {row[1] for row in connection.execute("PRAGMA table_info(documents)")}
        if "is_enabled" not in columns:
            connection.execute("ALTER TABLE documents ADD COLUMN is_enabled INTEGER NOT NULL DEFAULT 1")
        if "knowledge_base_id" not in columns:
            connection.execute("ALTER TABLE documents ADD COLUMN knowledge_base_id TEXT NOT NULL DEFAULT 'default'")
        if "category" not in columns:
            connection.execute("ALTER TABLE documents ADD COLUMN category TEXT NOT NULL DEFAULT 'general'")
        connection.execute("""
            CREATE TABLE IF NOT EXISTS knowledge_bases (
                id TEXT PRIMARY KEY, name TEXT NOT NULL, product_name TEXT NOT NULL,
                description TEXT NOT NULL DEFAULT '', is_enabled INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL
            )
        """)
        if connection.execute("SELECT 1 FROM knowledge_bases WHERE id = 'default'").fetchone() is None:
            connection.execute("INSERT INTO knowledge_bases VALUES (?, ?, ?, ?, ?, ?)",
                ("default", "通用知识库", "通用产品", "自动迁移的默认知识库", 1, datetime.now(UTC).isoformat()))
        if connection.execute("SELECT 1 FROM documents LIMIT 1").fetchone() is None and METADATA_PATH.exists():
            try:
                legacy_records = json.loads(METADATA_PATH.read_text(encoding="utf-8"))
            except json.JSONDecodeError:
                legacy_records = []
            for item in legacy_records:
                record = DocumentRecord.model_validate(item)
                _insert_record(connection, record)


def _document_connection() -> sqlite3.Connection:
    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row
    return connection


def _load_records() -> list[DocumentRecord]:
    with _METADATA_LOCK:
        _ensure_storage()
        with _document_connection() as connection:
            rows = connection.execute("SELECT * FROM documents ORDER BY created_at DESC").fetchall()
        return [DocumentRecord.model_validate(dict(row)) for row in rows]


def _save_records(records: list[DocumentRecord]) -> None:
    with _METADATA_LOCK:
        _ensure_storage()
        with _document_connection() as connection:
            connection.execute("DELETE FROM documents")
            for record in records:
                _insert_record(connection, record)


def _insert_record(connection: sqlite3.Connection, record: DocumentRecord) -> None:
    connection.execute(
        """
        INSERT OR REPLACE INTO documents (
            id, filename, content_type, status, size_bytes, text_length, created_at,
            error_message, index_status, indexed_chunks, index_error_message, is_enabled, knowledge_base_id, category
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            record.id,
            record.filename,
            record.content_type,
            record.status,
            record.size_bytes,
            record.text_length,
            record.created_at.isoformat(),
            record.error_message,
            record.index_status,
            record.indexed_chunks,
            record.index_error_message,
            int(record.is_enabled),
            record.knowledge_base_id,
            record.category,
        ),
    )


def _atomic_write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp",
            delete=False,
        ) as temporary_file:
            json.dump(payload, temporary_file, ensure_ascii=False, indent=2)
            temporary_file.flush()
            os.fsync(temporary_file.fileno())
            temporary_path = Path(temporary_file.name)
        os.replace(temporary_path, path)
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def _clean_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "".join(
        character
        for character in text
        if character in "\n\t" or (unicodedata.category(character) != "Cc" and character not in "\u200b\ufeff")
    )
    cleaned_lines = []
    for line in text.split("\n"):
        cells = [re.sub(r" {2,}", " ", cell.strip()) for cell in line.split("\t")]
        cleaned_lines.append("\t".join(cells).rstrip("\t"))
    return re.sub(r"\n{3,}", "\n\n", "\n".join(cleaned_lines)).strip()


def _extract_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    rendered_pdf = None
    page_texts: list[tuple[int, str]] = []
    try:
        for index, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            if not page_text.strip() and index <= MAX_OCR_PAGES:
                if rendered_pdf is None:
                    rendered_pdf = pymupdf.open(path)
                page_text = _ocr_pdf_page(rendered_pdf[index - 1])
            if page_text.strip():
                page_texts.append((index, page_text.strip()))
    finally:
        if rendered_pdf is not None:
            rendered_pdf.close()
    page_texts = _remove_repeated_page_margins(page_texts)
    return "\n\n".join(f"[page {index}]\n{text}" for index, text in page_texts)


def _ocr_pdf_page(page: pymupdf.Page) -> str:
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
    return _ocr_image_bytes(pixmap.tobytes("png"))


def _ocr_image_bytes(content: bytes) -> str:
    with _OCR_LOCK:
        global _OCR_ENGINE
        if _OCR_ENGINE is None:
            _OCR_ENGINE = RapidOCR()
        result, _ = _OCR_ENGINE(content)
    if not result:
        return ""
    return "\n".join(str(item[1]).strip() for item in result if len(item) > 1 and str(item[1]).strip())


def _remove_repeated_page_margins(page_texts: list[tuple[int, str]]) -> list[tuple[int, str]]:
    if len(page_texts) < 3:
        return page_texts
    edge_counts: Counter[str] = Counter()
    page_lines: list[tuple[int, list[str]]] = []
    for page_number, text in page_texts:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        page_lines.append((page_number, lines))
        page_edge_lines = {
            _margin_key(line)
            for line in lines[:2] + lines[-2:]
            if len(line.strip()) >= 3
        }
        edge_counts.update(page_edge_lines)
    threshold = max(2, int(len(page_texts) * 0.6 + 0.5))
    repeated = {line for line, count in edge_counts.items() if count >= threshold}
    if not repeated:
        return page_texts
    cleaned_pages = []
    for page_number, lines in page_lines:
        last_index = len(lines) - 1
        kept = [
            line for index, line in enumerate(lines)
            if not ((index < 2 or index >= last_index - 1) and _margin_key(line) in repeated)
        ]
        cleaned_pages.append((page_number, "\n".join(kept)))
    return cleaned_pages


def _margin_key(line: str) -> str:
    return re.sub(r"\d+", "#", re.sub(r"\s+", " ", line.strip().lower()))


def _extract_plain_text(path: Path) -> str:
    raw = path.read_bytes()
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        match = from_bytes(raw).best()
        if match is None:
            raise DocumentServiceError("无法识别文本文件编码")
        return str(match)


def _extract_docx_text(path: Path) -> str:
    document = DocxDocument(path)
    blocks: list[str] = []
    table_index = 0
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            style_name = block.style.name.lower() if block.style else ""
            heading_match = re.match(r"heading\s+(\d+)", style_name)
            if heading_match:
                text = f"{'#' * min(int(heading_match.group(1)), 6)} {text}"
            blocks.append(text)
        elif isinstance(block, Table):
            table_index += 1
            rows = []
            for row in block.rows:
                values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(values):
                    rows.append("\t".join(values))
            if rows:
                blocks.append(f"[table {table_index}]\n" + "\n".join(rows))
    return "\n\n".join(blocks)


def _format_cell_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _extract_xlsx_text(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            rows = []
            for row in worksheet.iter_rows(values_only=True):
                values = [_format_cell_value(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    rows.append("\t".join(values))
            if rows:
                sheets.append(f"[sheet {worksheet.title}]\n" + "\n".join(rows))
    finally:
        workbook.close()
    return "\n\n".join(sheets)


def _extract_xls_text(path: Path) -> str:
    workbook = xlrd.open_workbook(path, on_demand=True)
    sheets = []
    try:
        for worksheet in workbook.sheets():
            rows = []
            for row_index in range(worksheet.nrows):
                values = [_format_cell_value(value) for value in worksheet.row_values(row_index)]
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    rows.append("\t".join(values))
            if rows:
                sheets.append(f"[sheet {worksheet.name}]\n" + "\n".join(rows))
    finally:
        workbook.release_resources()
    return "\n\n".join(sheets)


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    _validate_file_signature(path, suffix)
    if suffix == ".pdf":
        return _clean_text(_extract_pdf_text(path))
    if suffix in {".txt", ".md", ".markdown", ".csv"}:
        return _clean_text(_extract_plain_text(path))
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
        return _clean_text(f"[image]\n{_ocr_image_bytes(path.read_bytes())}")
    if suffix == ".docx":
        return _clean_text(_extract_docx_text(path))
    if suffix == ".xlsx":
        return _clean_text(_extract_xlsx_text(path))
    if suffix == ".xls":
        return _clean_text(_extract_xls_text(path))
    raise DocumentServiceError("仅支持 PDF、图片、TXT、Markdown、CSV、Word、Excel 文档")


def _validate_file_signature(path: Path, suffix: str) -> None:
    prefix = path.read_bytes()[:8]
    if suffix == ".pdf" and not prefix.startswith(b"%PDF"):
        raise DocumentServiceError("文件扩展名为 PDF，但内容不是有效 PDF")
    if suffix in {".docx", ".xlsx"} and not prefix.startswith(b"PK"):
        raise DocumentServiceError("Office 文件内容无效；如果来自 Git LFS，请下载真实文件而不是指针文件")
    if suffix == ".xls" and prefix != bytes.fromhex("D0CF11E0A1B11AE1"):
        raise DocumentServiceError("文件扩展名为 XLS，但内容不是有效的旧版 Excel 工作簿")


async def save_uploaded_document(file: UploadFile, knowledge_base_id: str = "default", category: str = "general") -> DocumentRecord:
    _ensure_storage()
    if not get_knowledge_base(knowledge_base_id).is_enabled:
        raise DocumentServiceError("知识库已停用，不能上传文档")
    original_name = Path(file.filename or "").name
    suffix = Path(original_name).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise DocumentServiceError("仅支持 PDF、图片、TXT、Markdown、CSV、Word、Excel 文档")

    document_id = uuid4().hex
    stored_name = f"{document_id}{suffix}"
    upload_path = UPLOAD_DIR / stored_name
    text_path = TEXT_DIR / f"{document_id}.txt"

    size_bytes = 0
    with upload_path.open("wb") as output:
        while chunk := await file.read(1024 * 1024):
            size_bytes += len(chunk)
            if size_bytes > MAX_UPLOAD_SIZE_BYTES:
                upload_path.unlink(missing_ok=True)
                raise DocumentServiceError("单个文档不能超过 15MB")
            output.write(chunk)

    status = "processed"
    error_message = None
    text_length = 0

    try:
        parsed_text = await asyncio.to_thread(_extract_text, upload_path)
        text_length = len(parsed_text)
        text_path.write_text(parsed_text, encoding="utf-8")
        if not parsed_text:
            status = "failed"
            error_message = "文档未解析出有效文本"
    except Exception as exc:
        status = "failed"
        error_message = str(exc)
        text_path.write_text("", encoding="utf-8")

    record = DocumentRecord(
        id=document_id,
        filename=original_name,
        content_type=file.content_type or "application/octet-stream",
        status=status,
        size_bytes=size_bytes,
        text_length=text_length,
        created_at=datetime.now(UTC),
        error_message=error_message,
        knowledge_base_id=knowledge_base_id,
        category=category.strip()[:80] or "general",
    )

    with _METADATA_LOCK:
        records = _load_records()
        records.insert(0, record)
        _save_records(records)

        if record.status == "processed":
            from app.services.rag_service import index_document

            try:
                indexed_chunks = await asyncio.to_thread(index_document, record)
                record = update_document_index_status(
                    record.id,
                    index_status="indexed" if indexed_chunks else "failed",
                    indexed_chunks=indexed_chunks,
                    index_error_message=None if indexed_chunks else "文档未生成可索引片段",
                )
            except Exception as exc:
                record = update_document_index_status(
                    record.id,
                    index_status="failed",
                    indexed_chunks=0,
                    index_error_message=f"文档解析成功，但索引失败：{exc}",
                )

    return record


def list_documents() -> list[DocumentRecord]:
    return _load_records()


def list_knowledge_bases() -> list[KnowledgeBaseRecord]:
    _ensure_storage()
    with _document_connection() as connection:
        rows = connection.execute("SELECT * FROM knowledge_bases ORDER BY created_at DESC").fetchall()
    return [KnowledgeBaseRecord.model_validate(dict(row)) for row in rows]


def get_knowledge_base(knowledge_base_id: str) -> KnowledgeBaseRecord:
    _ensure_storage()
    with _document_connection() as connection:
        row = connection.execute("SELECT * FROM knowledge_bases WHERE id = ?", (knowledge_base_id,)).fetchone()
    if row is None:
        raise DocumentServiceError("知识库不存在")
    return KnowledgeBaseRecord.model_validate(dict(row))


def create_knowledge_base(request: KnowledgeBaseCreate) -> KnowledgeBaseRecord:
    _ensure_storage()
    record = KnowledgeBaseRecord(id=uuid4().hex, **request.model_dump(), created_at=datetime.now(UTC))
    with _document_connection() as connection:
        connection.execute("INSERT INTO knowledge_bases VALUES (?, ?, ?, ?, ?, ?)",
            (record.id, record.name, record.product_name, record.description, int(record.is_enabled), record.created_at.isoformat()))
    return record


def update_knowledge_base(knowledge_base_id: str, request: KnowledgeBaseUpdate) -> KnowledgeBaseRecord:
    current = get_knowledge_base(knowledge_base_id)
    updated = current.model_copy(update=request.model_dump(exclude_none=True))
    with _document_connection() as connection:
        connection.execute("UPDATE knowledge_bases SET name=?, product_name=?, description=?, is_enabled=? WHERE id=?",
            (updated.name, updated.product_name, updated.description, int(updated.is_enabled), updated.id))
    return updated


def delete_knowledge_base(knowledge_base_id: str) -> None:
    if knowledge_base_id == "default":
        raise DocumentServiceError("默认知识库不能删除")
    if any(record.knowledge_base_id == knowledge_base_id for record in list_documents()):
        raise DocumentServiceError("知识库仍有文档，请先迁移或删除文档")
    with _document_connection() as connection:
        if connection.execute("DELETE FROM knowledge_bases WHERE id=?", (knowledge_base_id,)).rowcount == 0:
            raise DocumentServiceError("知识库不存在")


def set_document_enabled(document_id: str, is_enabled: bool) -> DocumentRecord:
    """Disable retrieval without deleting the source or its vector index."""
    with _METADATA_LOCK:
        record = get_document(document_id)
        updated = record.model_copy(update={"is_enabled": is_enabled})
        _replace_record(updated)
        return updated


def update_document_index_status(
    document_id: str,
    *,
    index_status: DocumentIndexStatus,
    indexed_chunks: int = 0,
    index_error_message: str | None = None,
) -> DocumentRecord:
    with _METADATA_LOCK:
        records = _load_records()
        for index, record in enumerate(records):
            if record.id != document_id:
                continue
            updated_record = record.model_copy(
                update={
                    "index_status": index_status,
                    "indexed_chunks": indexed_chunks,
                    "index_error_message": index_error_message,
                }
            )
            records[index] = updated_record
            _save_records(records)
            return updated_record
        raise DocumentServiceError("文档不存在")


def get_document(document_id: str) -> DocumentRecord:
    for record in _load_records():
        if record.id == document_id:
            return record
    raise DocumentServiceError("文档不存在")


def get_document_preview(document_id: str, limit: int = 1200) -> tuple[DocumentRecord, str]:
    record = get_document(document_id)
    text_path = TEXT_DIR / f"{document_id}.txt"
    preview = text_path.read_text(encoding="utf-8", errors="ignore")[:limit] if text_path.exists() else ""
    return record, preview


def reprocess_document(document_id: str) -> DocumentRecord:
    record = get_document(document_id)
    upload_path = next((path for path in UPLOAD_DIR.glob(f"{document_id}.*") if path.is_file()), None)
    if upload_path is None:
        raise DocumentServiceError("原始文档文件不存在")

    text_path = TEXT_DIR / f"{document_id}.txt"
    try:
        parsed_text = _extract_text(upload_path)
        if not parsed_text:
            raise DocumentServiceError("文档未解析出有效文本")
        text_path.write_text(parsed_text, encoding="utf-8")
        record = record.model_copy(
            update={
                "status": "processed",
                "text_length": len(parsed_text),
                "error_message": None,
                "index_status": "not_indexed",
                "indexed_chunks": 0,
                "index_error_message": None,
            }
        )
        _replace_record(record)
    except DocumentServiceError:
        raise
    except Exception as exc:
        raise DocumentServiceError(f"文档重新解析失败：{exc}") from exc

    from app.services.rag_service import index_document

    try:
        indexed_chunks = index_document(record)
        return update_document_index_status(
            document_id,
            index_status="indexed" if indexed_chunks else "failed",
            indexed_chunks=indexed_chunks,
            index_error_message=None if indexed_chunks else "文档未生成可索引片段",
        )
    except Exception as exc:
        return update_document_index_status(
            document_id,
            index_status="failed",
            indexed_chunks=0,
            index_error_message=f"文档解析成功，但索引失败：{exc}",
        )


def _replace_record(updated_record: DocumentRecord) -> None:
    with _METADATA_LOCK:
        records = _load_records()
        for index, record in enumerate(records):
            if record.id == updated_record.id:
                records[index] = updated_record
                _save_records(records)
                return
        raise DocumentServiceError("文档不存在")


def delete_document(document_id: str) -> None:
    with _METADATA_LOCK:
        records = _load_records()
        next_records = [record for record in records if record.id != document_id]
        if len(next_records) == len(records):
            raise DocumentServiceError("文档不存在")

        from app.services.rag_service import delete_document_index

        delete_document_index(document_id)
        for path in UPLOAD_DIR.glob(f"{document_id}.*"):
            if path.is_file():
                path.unlink(missing_ok=True)
        text_path = TEXT_DIR / f"{document_id}.txt"
        text_path.unlink(missing_ok=True)
        _save_records(next_records)


def reset_documents_for_tests() -> None:
    if DATA_DIR.exists():
        shutil.rmtree(DATA_DIR)
